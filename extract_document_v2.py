#!/usr/bin/env python3

import argparse
import hashlib
import json
import mimetypes
import sys
from pathlib import Path


SCHEMA_VERSION = "document-manifest-v2"
EXTRACTOR_VERSION = "phase-a-1"
OCR_TEXT_CHAR_THRESHOLD = 40


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def normalize_text(text: str) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.splitlines()]
    return "\n".join(lines).strip()


def bbox_to_list(rect):
    if rect is None:
        return None
    try:
        return [round(float(rect.x0), 3), round(float(rect.y0), 3),
                round(float(rect.x1), 3), round(float(rect.y1), 3)]
    except AttributeError:
        try:
            values = list(rect)
            if len(values) == 4:
                return [round(float(x), 3) for x in values]
        except Exception:
            pass
    return None


def safe_cell(value):
    if value is None:
        return ""
    return normalize_text(str(value))


def extract_pdf(path: Path):
    try:
        import pymupdf
    except ImportError as exc:
        raise RuntimeError("PyMuPDF is required for PDF extraction") from exc

    doc = pymupdf.open(path)
    units = []
    total_tables = 0
    total_images = 0
    total_text_chars = 0
    ocr_candidates = []

    for page_index, page in enumerate(doc):
        page_number = page_index + 1
        text = normalize_text(page.get_text("text"))
        total_text_chars += len(text)

        tables = []
        if hasattr(page, "find_tables"):
            try:
                finder = page.find_tables()
                raw_tables = getattr(finder, "tables", finder) or []
                for table_index, table in enumerate(raw_tables):
                    extracted = table.extract() or []
                    rows = [
                        [safe_cell(cell) for cell in row]
                        for row in extracted
                    ]
                    column_count = max((len(row) for row in rows), default=0)
                    tables.append({
                        "table_index": table_index,
                        "bbox": bbox_to_list(getattr(table, "bbox", None)),
                        "row_count": len(rows),
                        "column_count": column_count,
                        "rows": rows,
                    })
            except Exception as exc:
                tables.append({
                    "table_index": 0,
                    "bbox": None,
                    "row_count": 0,
                    "column_count": 0,
                    "rows": [],
                    "extraction_error": f"{type(exc).__name__}: {exc}",
                })

        images = []
        try:
            image_info = page.get_image_info(xrefs=True) or []
        except TypeError:
            image_info = page.get_image_info() or []
        except Exception:
            image_info = []

        seen = set()
        for asset_index, info in enumerate(image_info):
            xref = info.get("xref")
            bbox = bbox_to_list(info.get("bbox"))
            dedupe_key = (xref, tuple(bbox) if bbox else None)
            if dedupe_key in seen:
                continue
            seen.add(dedupe_key)

            images.append({
                "asset_index": len(images),
                "xref": xref if isinstance(xref, int) and xref > 0 else None,
                "bbox": bbox,
                "width": info.get("width"),
                "height": info.get("height"),
                "colorspace": info.get("cs-name") or info.get("colorspace"),
                "bits_per_component": info.get("bpc"),
                "transform": list(info.get("transform")) if info.get("transform") else None,
            })

        if not images:
            try:
                for img in page.get_images(full=True):
                    xref = img[0] if img else None
                    rects = []
                    if xref:
                        try:
                            rects = page.get_image_rects(xref)
                        except Exception:
                            rects = []
                    if not rects:
                        rects = [None]
                    for rect in rects:
                        images.append({
                            "asset_index": len(images),
                            "xref": xref,
                            "bbox": bbox_to_list(rect),
                            "width": img[2] if len(img) > 2 else None,
                            "height": img[3] if len(img) > 3 else None,
                            "colorspace": None,
                            "bits_per_component": img[4] if len(img) > 4 else None,
                            "transform": None,
                        })
            except Exception:
                pass

        total_tables += len(tables)
        total_images += len(images)

        ocr_candidate = len(text.strip()) < OCR_TEXT_CHAR_THRESHOLD and len(images) > 0
        if ocr_candidate:
            ocr_candidates.append(page_number)

        units.append({
            "unit_index": page_index,
            "unit_type": "page",
            "page_number": page_number,
            "width": round(float(page.rect.width), 3),
            "height": round(float(page.rect.height), 3),
            "text": text,
            "text_chars": len(text),
            "tables": tables,
            "images": images,
            "ocr_candidate": ocr_candidate,
        })

    return {
        "document_kind": "pdf",
        "page_count": len(doc),
        "units": units,
        "summary": {
            "unit_count": len(units),
            "total_text_chars": total_text_chars,
            "total_tables": total_tables,
            "total_images": total_images,
            "ocr_candidate_units": ocr_candidates,
        },
        "extraction": {
            "text_extractor": "pymupdf",
            "table_extractor": "pymupdf.find_tables",
            "image_extractor": "pymupdf.image-info",
        },
    }


def extract_docx(path: Path):
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for DOCX extraction") from exc

    doc = Document(path)

    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = normalize_text(paragraph.text)
        if text:
            paragraphs.append({
                "paragraph_index": index,
                "style": paragraph.style.name if paragraph.style else None,
                "text": text,
            })

    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([
                normalize_text(cell.text)
                for cell in row.cells
            ])
        column_count = max((len(row) for row in rows), default=0)
        tables.append({
            "table_index": table_index,
            "bbox": None,
            "row_count": len(rows),
            "column_count": column_count,
            "rows": rows,
        })

    images = []
    rel_seen = set()
    for rel_id, rel in doc.part.rels.items():
        target = getattr(rel, "target_part", None)
        content_type = getattr(target, "content_type", None)
        if content_type and content_type.startswith("image/"):
            if rel_id in rel_seen:
                continue
            rel_seen.add(rel_id)
            partname = str(getattr(target, "partname", ""))
            images.append({
                "asset_index": len(images),
                "relationship_id": rel_id,
                "partname": partname or None,
                "content_type": content_type,
                "bbox": None,
                "width": None,
                "height": None,
            })

    text = "\n\n".join(item["text"] for item in paragraphs)
    ocr_candidate = len(text.strip()) < OCR_TEXT_CHAR_THRESHOLD and len(images) > 0

    unit = {
        "unit_index": 0,
        "unit_type": "document",
        "page_number": None,
        "text": text,
        "text_chars": len(text),
        "paragraphs": paragraphs,
        "tables": tables,
        "images": images,
        "ocr_candidate": ocr_candidate,
    }

    return {
        "document_kind": "docx",
        "page_count": None,
        "units": [unit],
        "summary": {
            "unit_count": 1,
            "total_text_chars": len(text),
            "total_tables": len(tables),
            "total_images": len(images),
            "ocr_candidate_units": [0] if ocr_candidate else [],
        },
        "extraction": {
            "text_extractor": "python-docx",
            "table_extractor": "python-docx",
            "image_extractor": "python-docx.relationships",
        },
    }


def extract_plain_text(path: Path):
    text = normalize_text(
        path.read_text(encoding="utf-8", errors="replace")
    )
    unit = {
        "unit_index": 0,
        "unit_type": "document",
        "page_number": None,
        "text": text,
        "text_chars": len(text),
        "tables": [],
        "images": [],
        "ocr_candidate": False,
    }

    return {
        "document_kind": path.suffix.lower().lstrip("."),
        "page_count": None,
        "units": [unit],
        "summary": {
            "unit_count": 1,
            "total_text_chars": len(text),
            "total_tables": 0,
            "total_images": 0,
            "ocr_candidate_units": [],
        },
        "extraction": {
            "text_extractor": "python.text",
            "table_extractor": None,
            "image_extractor": None,
        },
    }


def build_manifest(path: Path):
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        extracted = extract_pdf(path)
    elif suffix == ".docx":
        extracted = extract_docx(path)
    elif suffix in {".txt", ".md"}:
        extracted = extract_plain_text(path)
    else:
        raise RuntimeError(f"Unsupported file type: {suffix or '<none>'}")

    mime_type, _ = mimetypes.guess_type(path.name)

    return {
        "schema_version": SCHEMA_VERSION,
        "extractor_version": EXTRACTOR_VERSION,
        "source": {
            "filename": path.name,
            "suffix": suffix,
            "mime_type": mime_type,
            "size_bytes": path.stat().st_size,
            "sha256": sha256_file(path),
        },
        "document": {
            "kind": extracted["document_kind"],
            "page_count": extracted["page_count"],
        },
        "units": extracted["units"],
        "summary": extracted["summary"],
        "extraction": {
            **extracted["extraction"],
            "ocr_performed": False,
            "vision_model_used": False,
            "database_write_performed": False,
        },
    }


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Phase A canonical document extractor. "
            "Reads a file and emits a structured JSON manifest. "
            "No OCR, AI calls, embeddings, or database writes are performed."
        )
    )
    parser.add_argument(
        "file",
        help="PDF, DOCX, TXT or MD file",
    )
    parser.add_argument(
        "-o",
        "--output",
        help="Write JSON manifest to this path instead of stdout",
    )
    parser.add_argument(
        "--compact",
        action="store_true",
        help="Emit compact JSON instead of pretty-printed JSON",
    )
    args = parser.parse_args()

    path = Path(args.file)

    if not path.exists():
        raise SystemExit(f"File not found: {path}")
    if not path.is_file():
        raise SystemExit(f"Not a file: {path}")

    try:
        manifest = build_manifest(path)
    except Exception as exc:
        raise SystemExit(
            f"Extraction failed: {type(exc).__name__}: {exc}"
        ) from exc

    if args.compact:
        payload = json.dumps(
            manifest,
            ensure_ascii=False,
            separators=(",", ":"),
        )
    else:
        payload = json.dumps(
            manifest,
            indent=2,
            ensure_ascii=False,
        )

    if args.output:
        output = Path(args.output)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(payload + "\n", encoding="utf-8")
        print(output)
    else:
        print(payload)


if __name__ == "__main__":
    main()
